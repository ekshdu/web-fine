import datetime
import io
import os

import docx
from docxcompose.composer import Composer
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from config import settings

MALE_GENDER_ID = 1


def _money(value) -> str:
    if value is None:
        return ""
    text = f"{float(value):.2f}"
    return text.rstrip("0").rstrip(".")


def _date(value) -> str:
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.strftime("%d.%m.%Y")
    return str(value or "")


def _render_fine_document(rec: dict) -> docx.Document:
    """Заполненное постановление по шаблону (shablon1 — для id_gender=1, иначе shablon2)."""
    date_time = rec["date_time"]
    actual_speed = rec.get("actual_speed")
    limit_speed = rec.get("fine_speed")
    over_speed = actual_speed - limit_speed if (actual_speed and limit_speed and actual_speed > limit_speed) else 0
    half_sum = round(float(rec["fine_sum"]) / 2, 2) if rec["fine_sum"] else 0
    driver_fio = f"{rec['drv_surname']} {rec['drv_name']} {rec.get('drv_middle') or ''}".strip()

    replacements = {
        "{{ id }}": str(rec["num_post"]),
        "{{ date }}": date_time.strftime("%d.%m.%Y"),
        "{{ time }}": date_time.strftime("%H:%M"),
        "{{ address }}": str(rec["fine_place"]),
        "{{ car_mark }}": str(rec.get("car_model") or ""),
        "{{ car_num }}": str(rec["car_num"]),
        "{{ actual_speed_fine }}": str(actual_speed or ""),
        "{{ speed_fine }}": str(limit_speed or ""),
        "{{ up_speed_car }}": str(over_speed),
        "{{ car_id }}": str(rec["car_num"]),
        "{{ FIO_driver }}": driver_fio,
        "{{ date_of_birth }}": _date(rec.get("drv_dob")),
        "{{ place_of_birth }}": str(rec.get("drv_birth") or ""),
        "{{ section_num }}": str(rec.get("section_num") or ""),
        "{{ fine_sum }}": _money(rec["fine_sum"]),
        "{{ half_of_fine_sum }}": _money(half_sum),
        "{{ employee }}": str(rec.get("emp_fio") or "").strip(),
    }

    template_name = "shablon1.docx" if rec.get("id_gender") == MALE_GENDER_ID else "shablon2.docx"
    template_path = os.path.join(settings.templates_dir, template_name)
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Не найден шаблон {template_name}")

    document = docx.Document(template_path)

    def replace_in_paragraph(paragraph):
        full_text = "".join(run.text for run in paragraph.runs)
        changed = False
        for key, value in replacements.items():
            if key in full_text:
                full_text = full_text.replace(key, value)
                changed = True
        if changed and paragraph.runs:
            paragraph.runs[0].text = full_text
            for run in paragraph.runs[1:]:
                run.text = ""

    for paragraph in document.paragraphs:
        replace_in_paragraph(paragraph)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph)

    return document


def build_fine_word(rec: dict) -> bytes:
    buffer = io.BytesIO()
    _render_fine_document(rec).save(buffer)
    return buffer.getvalue()


def build_combined_word(records: list[dict]) -> bytes:
    """Один Word-документ со всеми постановлениями, каждое с новой страницы."""
    master = _render_fine_document(records[0])
    composer = Composer(master)
    for record in records[1:]:
        master.add_page_break()
        composer.append(_render_fine_document(record))
    buffer = io.BytesIO()
    composer.save(buffer)
    return buffer.getvalue()


def build_period_excel(rows: list[dict]) -> bytes:
    """Excel-реестр неоплаченных и просроченных постановлений за период."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Постановления"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    headers = ["ID", "Дата", "Постановление", "ФИО", "Госномер",
               "Место", "Статья", "Факт / Лимит ск.", "Сумма", "Статус"]
    widths = [6, 18, 20, 24, 12, 30, 16, 16, 10, 14]
    for index, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = Font(name="Arial", bold=True, size=11)
        cell.alignment = center
        cell.border = border
        sheet.column_dimensions[cell.column_letter].width = width

    for row_index, r in enumerate(rows, start=2):
        fio = f"{r['driver_surname']} {r['driver_name']} {r['driver_middle_name'] or ''}".strip()
        values = [
            r["id"],
            r["date_time"].strftime("%d.%m.%Y %H:%M"),
            r["num_post"],
            fio,
            r["car_num"],
            r["fine_place"],
            r.get("section_num") or "-",
            f"{r['actual_speed'] or '-'} / {r['fine_speed'] or '-'}",
            float(r["fine_sum"]),
            r["status_name"],
        ]
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.font = Font(name="Arial", size=10)
            cell.border = border
            cell.alignment = left if col_index in (4, 6) else center

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()