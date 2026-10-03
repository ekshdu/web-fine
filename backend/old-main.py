import sys
import datetime
import docx
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os
import pymysql
import pymysql.cursors
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QLineEdit, QStackedWidget, QMessageBox,
    QFrame, QTableWidget, QTableWidgetItem, QTabWidget, QHeaderView,
    QDialog, QComboBox, QCalendarWidget, QDateEdit, QSizePolicy
)
from PyQt6.QtCore import Qt, QDate, QPoint
from PyQt6.QtGui import QFont
def setup_table(table: QTableWidget, font_size: int = 12):
    font = QFont("Arial", font_size)
    table.setFont(font)
    table.horizontalHeader().setFont(QFont("Arial", font_size, QFont.Weight.Bold))
    table.setWordWrap(True)
    table.verticalHeader().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
    table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
    table.setTextElideMode(Qt.TextElideMode.ElideNone)
class DatabaseManager:
    def __init__(self):
        self.connection = None
        self.connect_to_db()
    def connect_to_db(self):
        try:
            self.connection = pymysql.connect(
                host='127.0.0.1', user='root', password='1qwerty', port=3306,
                database='fine', cursorclass=pymysql.cursors.DictCursor, connect_timeout=10
            )
        except Exception as e:
            print(f"Ошибка БД: {e}")
            self.connection = None
    def auth_employee(self, username_text, password_text):
        if not self.connection:
            return None
        try:
            cursor = self.connection.cursor()
            cursor.execute(
                "SELECT e.*, u.username FROM emp_gibdd e "
                "JOIN user_emp u ON e.id_user = u.id "
                "WHERE u.username = %s AND u.password = %s",
                (username_text, password_text)
            )
            user = cursor.fetchone()
            cursor.close()
            return user
        except Exception as e:
            print(f"Ошибка авторизации: {e}")
            return None
    def get_driver(self, plate_number):
        if not self.connection:
            return None
        try:
            cursor = self.connection.cursor()
            cursor.execute(
                "SELECT d.id AS driver_id, d.name, d.surname, d.middle_name, "
                "d.place_of_birth, d.date_of_birth, d.id_gender, "
                "c.id AS car_id, c.car_num, c.car_model "
                "FROM car c JOIN driver d ON c.id_driver = d.id "
                "WHERE c.car_num = %s",
                (plate_number,)
            )
            res = cursor.fetchone()
            cursor.close()
            return res
        except Exception as e:
            print(f"Ошибка получения водителя: {e}")
            return None
    def reg_driver(self, surname, name, midname, plate, model, birth, dob, gender_id):
        if not self.connection:
            return False
        try:
            cursor = self.connection.cursor()
            cursor.execute(
                "INSERT INTO driver (name, surname, middle_name, place_of_birth, date_of_birth, id_gender) "
                "VALUES (%s, %s, %s, %s, %s, %s)",
                (name, surname, midname, birth, dob, gender_id)
            )
            driver_id = cursor.lastrowid
            cursor.execute(
                "INSERT INTO car (id_driver, car_num, car_model) VALUES (%s, %s, %s)",
                (driver_id, plate, model)
            )
            self.connection.commit()
            cursor.close()
            return True
        except Exception as e:
            self.connection.rollback()
            print(f"Ошибка регистрации: {e}")
            return False
    def get_cursor(self):
        try:
            self.connection.ping(reconnect=True)
        except Exception:
            self.connect_to_db()
        return self.connection.cursor()
    def update_overdue_fines(self):
        if not self.connection:
            return
        try:
            cursor = self.connection.cursor()
            cursor.execute(
                "UPDATE traffic_fine SET id_status_fine = 3 "
                "WHERE id_status_fine = 2 AND DATEDIFF(NOW(), date_time) >= 60"
            )
            self.connection.commit()
            cursor.close()
        except Exception as e:
            print(f"Ошибка обновления просроченных штрафов: {e}")
    def close_db(self):
        if self.connection and self.connection.open:
            self.connection.close()
class PeriodPickerPopup(QDialog):
    def __init__(self, parent_widget, on_period_selected):
        super().__init__(parent_widget, Qt.WindowType.Popup)
        self.on_period_selected = on_period_selected
        self.date_from = None
        self.selecting_second = False
        self.build_ui()
    def build_ui(self):
        self.setWindowFlags(Qt.WindowType.Popup | Qt.WindowType.FramelessWindowHint)
        self.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)
        self.setStyleSheet("PeriodPickerPopup { background: #ffffff; border: 1px solid #aab; border-radius: 6px; }")
        layout = QVBoxLayout(self)
        layout.setContentsMargins(8, 8, 8, 8)
        layout.setSpacing(5)
        self.lbl_hint = QLabel("Кликните на начальную дату")
        self.lbl_hint.setFont(QFont("Arial", 10))
        self.lbl_hint.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.lbl_hint)
        self.calendar = QCalendarWidget()
        self.calendar.setGridVisible(True)
        self.calendar.setMaximumDate(QDate.currentDate())
        self.calendar.setFixedSize(320, 220)
        self.calendar.clicked.connect(self.on_date_clicked)
        layout.addWidget(self.calendar)
        self.lbl_range = QLabel("—")
        self.lbl_range.setFont(QFont("Arial", 10))
        self.lbl_range.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_range.setStyleSheet("color: #2B5278; font-weight: bold;")
        layout.addWidget(self.lbl_range)
        btn_row = QHBoxLayout()
        btn_reset = QPushButton("Сбросить")
        btn_reset.setFixedHeight(30)
        btn_reset.setFont(QFont("Arial", 10))
        btn_reset.clicked.connect(self.reset_dates)
        btn_close = QPushButton("Закрыть")
        btn_close.setFixedHeight(30)
        btn_close.setFont(QFont("Arial", 10))
        btn_close.clicked.connect(self.close)
        btn_row.addWidget(btn_reset)
        btn_row.addWidget(btn_close)
        layout.addLayout(btn_row)
        self.adjustSize()
    def on_date_clicked(self, date):
        if not self.selecting_second:
            self.date_from = date
            self.selecting_second = True
            self.lbl_hint.setText("Теперь кликните конечную дату")
            self.lbl_range.setText(f"{date.toString('dd.MM.yyyy')}  —  ?")
        else:
            date_from = self.date_from
            date_to = date
            if date_to < date_from:
                date_from, date_to = date_to, date_from
            self.selecting_second = False
            self.lbl_hint.setText("Готово. Можно сбросить или закрыть.")
            self.lbl_range.setText(f"{date_from.toString('dd.MM.yyyy')}  —  {date_to.toString('dd.MM.yyyy')}")
            self.on_period_selected(date_from, date_to)
            self.close()
    def reset_dates(self):
        self.date_from = None
        self.selecting_second = False
        self.lbl_hint.setText("Кликните на начальную дату")
        self.lbl_range.setText("—")
        self.on_period_selected(None, None)
class RoleWindow(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.build_ui()
    def build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(60, 80, 60, 80)
        layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.setSpacing(20)
        lbl = QLabel("Система мониторинга штрафов ГИБДД")
        lbl.setFont(QFont("Arial", 20, QFont.Weight.Bold))
        layout.addWidget(lbl, alignment=Qt.AlignmentFlag.AlignHCenter)
        btn_drv = QPushButton("Вход по госномеру (Водитель)")
        btn_drv.setFixedSize(400, 50)
        btn_drv.setFont(QFont("Arial", 14))
        btn_drv.clicked.connect(lambda: self.main_window.stacked.setCurrentIndex(1))
        layout.addWidget(btn_drv, alignment=Qt.AlignmentFlag.AlignHCenter)
        btn_emp = QPushButton("Вход для сотрудника")
        btn_emp.setFixedSize(400, 50)
        btn_emp.setFont(QFont("Arial", 14))
        btn_emp.clicked.connect(lambda: self.main_window.stacked.setCurrentIndex(5))
        layout.addWidget(btn_emp, alignment=Qt.AlignmentFlag.AlignHCenter)
class DriverLoginWindow(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.build_ui()
    def build_ui(self):
        layout = QVBoxLayout(self)
        layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.setSpacing(15)
        lbl = QLabel("ВХОД ДЛЯ ВОДИТЕЛЯ")
        lbl.setFont(QFont("Arial", 20, QFont.Weight.Bold))
        layout.addWidget(lbl, alignment=Qt.AlignmentFlag.AlignHCenter)
        self.input_plate = QLineEdit()
        self.input_plate.setPlaceholderText("Введите госномер (А123АА77)")
        self.input_plate.setFixedSize(400, 50)
        self.input_plate.setFont(QFont("Arial", 18))
        self.input_plate.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(self.input_plate, alignment=Qt.AlignmentFlag.AlignHCenter)
        btn_login = QPushButton("Войти")
        btn_reg = QPushButton("Регистрация")
        btn_back = QPushButton("Назад")
        for btn in [btn_login, btn_reg, btn_back]:
            btn.setFixedSize(300, 45)
            btn.setFont(QFont("Arial", 14))
            layout.addWidget(btn, alignment=Qt.AlignmentFlag.AlignHCenter)
        btn_login.clicked.connect(self.check_login)
        btn_reg.clicked.connect(self.go_to_reg)
        btn_back.clicked.connect(lambda: self.main_window.stacked.setCurrentIndex(0))
    def go_to_reg(self):
        self.main_window.window_driver_reg.load_genders()
        self.main_window.stacked.setCurrentIndex(2)
    def check_login(self):
        plate = self.input_plate.text().strip().upper()
        if not plate:
            QMessageBox.warning(self, "Ошибка", "Введите госномер")
            return
        driver = self.main_window.db.get_driver(plate)
        if driver:
            self.input_plate.clear()
            self.main_window.driver_data = driver
            self.main_window.window_driver_main.load_data()
            self.main_window.stacked.setCurrentIndex(3)
        else:
            QMessageBox.information(self, "Не найдено", "Водитель не найден.")
class DriverRegWindow(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.build_ui()
    def build_ui(self):
        layout = QVBoxLayout(self)
        layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.setSpacing(10)
        lbl = QLabel("РЕГИСТРАЦИЯ")
        lbl.setFont(QFont("Arial", 20, QFont.Weight.Bold))
        layout.addWidget(lbl, alignment=Qt.AlignmentFlag.AlignHCenter)
        self.inp_sur = QLineEdit()
        self.inp_sur.setPlaceholderText("Фамилия")
        self.inp_name = QLineEdit()
        self.inp_name.setPlaceholderText("Имя")
        self.inp_mid = QLineEdit()
        self.inp_mid.setPlaceholderText("Отчество")
        self.combo_gender = QComboBox()
        self.combo_gender.setFixedSize(450, 45)
        self.combo_gender.setFont(QFont("Arial", 16))
        self.inp_plate = QLineEdit()
        self.inp_plate.setPlaceholderText("Госномер")
        self.inp_model = QLineEdit()
        self.inp_model.setPlaceholderText("Модель авто")
        self.inp_birth = QLineEdit()
        self.inp_birth.setPlaceholderText("Место рождения (город)")
        self.inp_dob = QLineEdit()
        self.inp_dob.setPlaceholderText("Год или дата рождения")
        for w in [self.inp_sur, self.inp_name, self.inp_mid, self.combo_gender,
                  self.inp_plate, self.inp_model, self.inp_birth, self.inp_dob]:
            w.setFixedSize(450, 45)
            w.setFont(QFont("Arial", 16))
            layout.addWidget(w, alignment=Qt.AlignmentFlag.AlignHCenter)
        btn_reg = QPushButton("Зарегистрироваться")
        btn_back = QPushButton("Отмена")
        for btn in [btn_reg, btn_back]:
            btn.setFixedSize(350, 45)
            btn.setFont(QFont("Arial", 14))
            layout.addWidget(btn, alignment=Qt.AlignmentFlag.AlignHCenter)
        btn_reg.clicked.connect(self.do_register)
        btn_back.clicked.connect(lambda: self.main_window.stacked.setCurrentIndex(1))
    def load_genders(self):
        self.combo_gender.clear()
        try:
            cur = self.main_window.db.get_cursor()
            cur.execute("SELECT id, name FROM gender ORDER BY id")
            for r in cur.fetchall():
                self.combo_gender.addItem(r['name'], r['id'])
            cur.close()
        except Exception as e:
            print("Ошибка загрузки пола:", e)
    def do_register(self):
        s = self.inp_sur.text().strip()
        n = self.inp_name.text().strip()
        m = self.inp_mid.text().strip()
        p = self.inp_plate.text().strip().upper()
        mo = self.inp_model.text().strip()
        b = self.inp_birth.text().strip()
        dob = self.inp_dob.text().strip()
        gender_id = self.combo_gender.currentData()
        if not (s and n and p and mo):
            QMessageBox.warning(self, "Ошибка", "Заполните обязательные поля (ФИО, авто)")
            return
        if not s.isalpha():
            QMessageBox.warning(self, "Ошибка", "Фамилия должна содержать только буквы")
            return
        if not n.isalpha():
            QMessageBox.warning(self, "Ошибка", "Имя должно содержать только буквы")
            return
        if m and not m.isalpha():
            QMessageBox.warning(self, "Ошибка", "Отчество должно содержать только буквы")
            return
        if b and not b.isalpha():
            QMessageBox.warning(self, "Ошибка", "Место рождения должно содержать только буквы")
            return
        if self.main_window.db.reg_driver(s, n, m, p, mo, b, dob, gender_id):
            QMessageBox.information(self, "Готово", "Регистрация прошла успешно")
            self.main_window.stacked.setCurrentIndex(1)
        else:
            QMessageBox.critical(self, "Ошибка", "Ошибка БД при регистрации.")
class DriverMainWindow(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.build_ui()
    def build_ui(self):
        layout = QVBoxLayout(self)
        nav = QHBoxLayout()
        btn_out = QPushButton("Выйти")
        btn_out.setFixedSize(150, 40)
        btn_out.clicked.connect(self.logout)
        nav.addWidget(btn_out)
        nav.addStretch()
        self.btn_profile = QPushButton("Профиль")
        self.btn_profile.setFixedSize(280, 40)
        self.btn_profile.setFont(QFont("Arial", 13))
        self.btn_profile.clicked.connect(self.open_profile)
        nav.addWidget(self.btn_profile)
        layout.addLayout(nav)
        btn_pay = QPushButton("Оплатить выбранный штраф")
        btn_pay.setFixedSize(350, 45)
        btn_pay.setFont(QFont("Arial", 14))
        btn_pay.clicked.connect(self.go_to_pay)
        layout.addWidget(btn_pay, alignment=Qt.AlignmentFlag.AlignLeft)
        self.table = QTableWidget()
        self.table.setColumnCount(7)
        self.table.setHorizontalHeaderLabels(
            ["Дата", "Постановление", "Госномер", "Модель", "Место", "Сумма", "Статус"]
        )
        setup_table(self.table, font_size=12)
        layout.addWidget(self.table)
    def load_data(self):
        self.main_window.db.update_overdue_fines()
        drv = self.main_window.driver_data
        fio = f"{drv['surname']} {drv['name']} {drv.get('middle_name', '') or ''}".strip()
        self.btn_profile.setText(f"{fio}")
        try:
            cur = self.main_window.db.get_cursor()
            cur.execute("""
                SELECT tf.id, tf.date_time, tf.num_post,
                       c.car_num, c.car_model, tf.fine_place,
                       fs.summary, sf.name AS status
                FROM traffic_fine tf
                JOIN car c ON tf.id_car = c.id
                JOIN fine_sum fs ON tf.id_fine_sum = fs.id
                JOIN status_fine sf ON tf.id_status_fine = sf.id
                WHERE c.id_driver = %s
                ORDER BY tf.date_time DESC
            """, (drv['driver_id'],))
            fines = cur.fetchall()
            cur.close()
            self.table.setRowCount(len(fines))
            for i, f in enumerate(fines):
                self.table.setItem(i, 0, QTableWidgetItem(str(f['date_time'])))
                self.table.setItem(i, 1, QTableWidgetItem(str(f['num_post'])))
                self.table.setItem(i, 2, QTableWidgetItem(str(f['car_num'])))
                self.table.setItem(i, 3, QTableWidgetItem(str(f['car_model'])))
                self.table.setItem(i, 4, QTableWidgetItem(str(f['fine_place'])))
                self.table.setItem(i, 5, QTableWidgetItem(str(f['summary'])))
                self.table.setItem(i, 6, QTableWidgetItem(str(f['status'])))
                self.table.item(i, 0).setData(Qt.ItemDataRole.UserRole, str(f['id']))
            self.table.resizeRowsToContents()
        except Exception as e:
            print(f"Ошибка таблицы водителя: {e}")
    def open_profile(self):
        DriverProfileDialog(self.main_window).exec()
    def go_to_pay(self):
        row = self.table.currentRow()
        if row == -1:
            QMessageBox.warning(self, "Внимание", "Выберите штраф для оплаты")
            return
        fine_data = {
            'id': self.table.item(row, 0).data(Qt.ItemDataRole.UserRole),
            'date_time': self.table.item(row, 0).text(),
            'num_post': self.table.item(row, 1).text(),
            'car_num': self.table.item(row, 2).text(),
            'car_model': self.table.item(row, 3).text(),
            'fine_place': self.table.item(row, 4).text(),
            'summary': self.table.item(row, 5).text(),
            'status': self.table.item(row, 6).text(),
        }
        self.main_window.window_driver_pay.load_fine(fine_data)
        self.main_window.stacked.setCurrentIndex(4)
    def logout(self):
        self.main_window.driver_data = None
        self.main_window.stacked.setCurrentIndex(0)
class DriverPayWindow(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.current_fine_id = None
        self.build_ui()
    def build_ui(self):
        layout = QVBoxLayout(self)
        layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl = QLabel("ОПЛАТА ШТРАФА")
        lbl.setFont(QFont("Arial", 20, QFont.Weight.Bold))
        layout.addWidget(lbl, alignment=Qt.AlignmentFlag.AlignHCenter)
        self.lbl_info = QLabel("Информация...")
        self.lbl_info.setFont(QFont("Arial", 16))
        self.lbl_info.setStyleSheet("background-color: #f0f0f0; border: 1px solid #ccc; padding: 20px; color: black;")
        self.lbl_info.setFixedSize(500, 300)
        self.lbl_info.setWordWrap(True)
        layout.addWidget(self.lbl_info, alignment=Qt.AlignmentFlag.AlignHCenter)
        layout.addSpacing(20)
        btn_pay = QPushButton("Оплатить")
        btn_cancel = QPushButton("Отмена")
        for btn in [btn_pay, btn_cancel]:
            btn.setFixedSize(300, 45)
            btn.setFont(QFont("Arial", 14))
            layout.addWidget(btn, alignment=Qt.AlignmentFlag.AlignHCenter)
        btn_pay.clicked.connect(self.process_payment)
        btn_cancel.clicked.connect(self.go_back)
    def load_fine(self, fine_data):
        self.current_fine_id = fine_data['id']
        self.lbl_info.setText(
            f"Дата: {fine_data['date_time']}\nПостановление: {fine_data['num_post']}\n"
            f"Госномер: {fine_data['car_num']}\nМесто: {fine_data['fine_place']}\n"
            f"Сумма: {fine_data['summary']} руб.\nСтатус: {fine_data['status']}"
        )
    def process_payment(self):
        try:
            cur = self.main_window.db.get_cursor()
            cur.execute("UPDATE traffic_fine SET id_status_fine = 1 WHERE id = %s", (self.current_fine_id,))
            self.main_window.db.connection.commit()
            cur.close()
            QMessageBox.information(self, "Готово", "Штраф успешно оплачен")
            self.go_back()
        except Exception as e:
            print(f"Ошибка оплаты: {e}")
    def go_back(self):
        self.main_window.window_driver_main.load_data()
        self.main_window.stacked.setCurrentIndex(3)
class DriverProfileDialog(QDialog):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.setWindowTitle("Настройки профиля")
        self.resize(500, 280)
        self.build_ui()
    def build_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(10)
        lbl = QLabel("Добавить автомобиль")
        lbl.setFont(QFont("Arial", 16, QFont.Weight.Bold))
        layout.addWidget(lbl, alignment=Qt.AlignmentFlag.AlignHCenter)
        drv = self.main_window.driver_data
        fio = f"{drv['surname']} {drv['name']} {drv.get('middle_name', '') or ''}".strip()
        lbl_fio = QLabel(f"Водитель: {fio}")
        lbl_fio.setFont(QFont("Arial", 13))
        layout.addWidget(lbl_fio, alignment=Qt.AlignmentFlag.AlignHCenter)
        self.inp_plate = QLineEdit()
        self.inp_plate.setPlaceholderText("Госномер новой машины")
        self.inp_model = QLineEdit()
        self.inp_model.setPlaceholderText("Модель новой машины")
        for w in [self.inp_plate, self.inp_model]:
            w.setFixedSize(420, 42)
            w.setFont(QFont("Arial", 14))
            layout.addWidget(w, alignment=Qt.AlignmentFlag.AlignHCenter)
        btn_add = QPushButton("Добавить автомобиль")
        btn_add.setFixedSize(300, 45)
        btn_add.setFont(QFont("Arial", 14))
        btn_add.clicked.connect(self.add_car)
        layout.addWidget(btn_add, alignment=Qt.AlignmentFlag.AlignHCenter)
        btn_close = QPushButton("Закрыть")
        btn_close.setFixedSize(200, 40)
        btn_close.setFont(QFont("Arial", 13))
        btn_close.clicked.connect(self.reject)
        layout.addWidget(btn_close, alignment=Qt.AlignmentFlag.AlignHCenter)
    def add_car(self):
        plate = self.inp_plate.text().strip().upper()
        model = self.inp_model.text().strip()
        if not plate or not model:
            QMessageBox.warning(self, "Ошибка", "Заполните госномер и модель")
            return
        driver_id = self.main_window.driver_data['driver_id']
        try:
            cur = self.main_window.db.get_cursor()
            cur.execute("SELECT id FROM car WHERE car_num = %s", (plate,))
            if cur.fetchone():
                QMessageBox.warning(self, "Ошибка", "Автомобиль с таким госномером уже зарегистрирован")
                cur.close()
                return
            cur.execute(
                "INSERT INTO car (id_driver, car_num, car_model) VALUES (%s, %s, %s)",
                (driver_id, plate, model)
            )
            self.main_window.db.connection.commit()
            cur.close()
            QMessageBox.information(self, "Готово", f"Автомобиль {plate} добавлен")
            self.inp_plate.clear()
            self.inp_model.clear()
        except Exception as e:
            self.main_window.db.connection.rollback()
            QMessageBox.critical(self, "Ошибка", f"Не удалось добавить:\n{e}")
class EmployeeLoginWindow(QWidget):
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.build_ui()
    def build_ui(self):
        layout = QVBoxLayout(self)
        layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.setSpacing(15)
        lbl = QLabel("ВХОД СОТРУДНИКА")
        lbl.setFont(QFont("Arial", 20, QFont.Weight.Bold))
        layout.addWidget(lbl, alignment=Qt.AlignmentFlag.AlignHCenter)
        self.inp_login = QLineEdit()
        self.inp_login.setPlaceholderText("Логин")
        self.inp_pass = QLineEdit()
        self.inp_pass.setPlaceholderText("Пароль")
        self.inp_pass.setEchoMode(QLineEdit.EchoMode.Password)
        for w in [self.inp_login, self.inp_pass]:
            w.setFixedSize(400, 50)
            w.setFont(QFont("Arial", 16))
            layout.addWidget(w, alignment=Qt.AlignmentFlag.AlignHCenter)
        btn_in = QPushButton("Войти")
        btn_back = QPushButton("Назад")
        for btn in [btn_in, btn_back]:
            btn.setFixedSize(300, 45)
            btn.setFont(QFont("Arial", 14))
            layout.addWidget(btn, alignment=Qt.AlignmentFlag.AlignHCenter)
        btn_in.clicked.connect(self.login_emp)
        btn_back.clicked.connect(lambda: self.main_window.stacked.setCurrentIndex(0))
    def login_emp(self):
        login = self.inp_login.text().strip()
        password = self.inp_pass.text().strip()
        if not login or not password:
            QMessageBox.warning(self, "Ошибка", "Введите логин и пароль")
            return
        try:
            user = self.main_window.db.auth_employee(login, password)
            if user:
                self.inp_login.clear()
                self.inp_pass.clear()
                self.main_window.emp_data = user
                self.main_window.window_emp_main.load_all_tabs()
                self.main_window.stacked.setCurrentIndex(6)
            else:
                QMessageBox.critical(self, "Ошибка", "Неверный логин/пароль")
        except Exception as e:
            QMessageBox.critical(self, "Критическая ошибка", str(e))
class EmployeeMainWindow(QWidget):
    BASE_QUERY = (
        "SELECT tf.id, tf.date_time, tf.num_post, "
        "TRIM(CONCAT(d.surname, ' ', d.name, ' ', COALESCE(d.middle_name, ''))) AS fio, "
        "c.car_num, tf.fine_place, "
        "fsec.section_num, tf.actual_speed, fsp.fine_speedcol, "
        "fs.summary, sf.name AS status "
        "FROM traffic_fine tf "
        "JOIN car c ON tf.id_car = c.id "
        "JOIN driver d ON c.id_driver = d.id "
        "JOIN fine_sum fs ON tf.id_fine_sum = fs.id "
        "JOIN status_fine sf ON tf.id_status_fine = sf.id "
        "LEFT JOIN fine_section fsec ON tf.id_fine_section = fsec.id "
        "LEFT JOIN fine_speed fsp ON tf.id_speed_fine = fsp.id"
    )
    def __init__(self, main_window):
        super().__init__()
        self.main_window = main_window
        self.report_date_from = None
        self.report_date_to = None
        self.build_ui()
    def build_ui(self):
        layout = QVBoxLayout(self)
        top_lay = QHBoxLayout()
        top_lay.addStretch()
        btn_out = QPushButton("Выход")
        btn_out.setFixedSize(150, 40)
        btn_out.clicked.connect(self.logout)
        top_lay.addWidget(btn_out)
        layout.addLayout(top_lay)
        self.tabs = QTabWidget()
        self.tabs.setStyleSheet("QTabBar::tab { font-size: 16pt; padding: 8px 24px; }")
        layout.addWidget(self.tabs)
        self.tab_svodka = QWidget()
        self.build_svodka()
        self.tabs.addTab(self.tab_svodka, "Сводка")
        self.tab_fines = QWidget()
        self.build_fines()
        self.tabs.addTab(self.tab_fines, "Штрафы")
        self.tab_report = QWidget()
        self.build_report()
        self.tabs.addTab(self.tab_report, "Отчет")
    headers = [
        "ID", "Дата", "Постановление", "ФИО", "Госномер",
        "Место", "Статья", "Факт/Лимит ск.", "Сумма", "Статус"
    ]
    def build_svodka(self):
        layout = QVBoxLayout(self.tab_svodka)
        layout.setSpacing(10)
        stat_lay = QHBoxLayout()
        stat_lay.setSpacing(12)
        self.lbl_all = QLabel("Всего штрафов:\n0")
        self.lbl_paid = QLabel("Оплаченных:\n0")
        self.lbl_unpaid = QLabel("Не оплачено\n0")
        self.lbl_debt = QLabel("Задолженность\n0")
        self.lbl_sum = QLabel("Общий долг:\n0 руб.")
        for lbl in [self.lbl_all, self.lbl_paid, self.lbl_unpaid, self.lbl_debt, self.lbl_sum]:
            lbl.setFrameStyle(QFrame.Shape.Box)
            lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
            lbl.setFont(QFont("Arial", 14))
            lbl.setFixedHeight(110)
            lbl.setWordWrap(True)
            stat_lay.addWidget(lbl)
        layout.addLayout(stat_lay)
        self.table_svodka = QTableWidget()
        self.table_svodka.setColumnCount(len(self.headers))
        self.table_svodka.setHorizontalHeaderLabels(self.headers)
        setup_table(self.table_svodka, font_size=12)
        layout.addWidget(self.table_svodka)
    def build_fines(self):
        layout = QVBoxLayout(self.tab_fines)
        top_lay = QHBoxLayout()
        self.inp_search = QLineEdit()
        self.inp_search.setPlaceholderText("Поиск по номеру постановления...")
        self.inp_search.setFixedSize(300, 45)
        self.inp_search.setFont(QFont("Arial", 14))
        btn_search = QPushButton("Найти")
        btn_search.setFixedSize(100, 45)
        btn_search.clicked.connect(self.load_fines_table)
        btn_add = QPushButton("Добавить")
        btn_edit = QPushButton("Редактировать")
        btn_del = QPushButton("Удалить")
        top_lay.addWidget(self.inp_search)
        top_lay.addWidget(btn_search)
        for btn in [btn_add, btn_edit, btn_del]:
            btn.setFixedHeight(45)
            btn.setFont(QFont("Arial", 13))
            btn.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
            top_lay.addWidget(btn)
        btn_add.clicked.connect(self.add_fine)
        btn_edit.clicked.connect(self.edit_fine)
        btn_del.clicked.connect(self.del_fine)
        layout.addLayout(top_lay)
        self.table_fines = QTableWidget()
        self.table_fines.setColumnCount(len(self.headers))
        self.table_fines.setHorizontalHeaderLabels(self.headers)
        setup_table(self.table_fines, font_size=12)
        layout.addWidget(self.table_fines)
    def build_report(self):
        layout = QVBoxLayout(self.tab_report)
        top_lay = QHBoxLayout()
        top_lay.setSpacing(8)
        self.btn_period = QPushButton("Период")
        self.btn_period.setFixedSize(90, 36)
        self.btn_period.setFont(QFont("Arial", 12))
        self.btn_period.setToolTip("Выбрать период")
        self.btn_period.clicked.connect(self.open_period_popup)
        self.lbl_period = QLabel("Период не выбран")
        self.lbl_period.setFont(QFont("Arial", 12))
        self.lbl_period.setFixedHeight(36)
        btn_reset = QPushButton("Сбросить")
        btn_reset.setFixedSize(90, 36)
        btn_reset.setFont(QFont("Arial", 12))
        btn_reset.setToolTip("Сбросить период")
        btn_reset.clicked.connect(self.reset_period)
        self.inp_report_post = QLineEdit()
        self.inp_report_post.setPlaceholderText("Поиск по постановлению")
        self.inp_report_post.setFixedSize(220, 36)
        self.inp_report_post.setFont(QFont("Arial", 12))
        btn_search_rep = QPushButton("Найти")
        btn_search_rep.setFixedSize(90, 36)
        btn_search_rep.setFont(QFont("Arial", 12))
        btn_search_rep.clicked.connect(self.load_report_table)
        btn_word = QPushButton("Word")
        btn_word.setFixedSize(110, 36)
        btn_word.setFont(QFont("Arial", 12))
        btn_word.clicked.connect(self.export_word)
        btn_excel = QPushButton("Excel")
        btn_excel.setFixedSize(110, 36)
        btn_excel.setFont(QFont("Arial", 12))
        btn_excel.clicked.connect(self.export_excel)
        for w in [self.btn_period, self.lbl_period, btn_reset,
                  self.inp_report_post, btn_search_rep, btn_word, btn_excel]:
            top_lay.addWidget(w)
        top_lay.addStretch()
        layout.addLayout(top_lay)
        self.table_report = QTableWidget()
        self.table_report.setColumnCount(len(self.headers))
        self.table_report.setHorizontalHeaderLabels(self.headers)
        setup_table(self.table_report, font_size=12)
        layout.addWidget(self.table_report)
    def open_period_popup(self):
        popup = PeriodPickerPopup(self, self.on_period_selected)
        btn_pos = self.btn_period.mapToGlobal(QPoint(0, self.btn_period.height() + 2))
        popup.move(btn_pos)
        popup.exec()
    def on_period_selected(self, date_from, date_to):
        self.report_date_from = date_from
        self.report_date_to = date_to
        if date_from and date_to:
            self.lbl_period.setText(
                f"{date_from.toString('dd.MM.yyyy')}  —  {date_to.toString('dd.MM.yyyy')}"
            )
        else:
            self.lbl_period.setText("Период не выбран")
        self.load_report_table()
    def reset_period(self):
        self.report_date_from = None
        self.report_date_to = None
        self.lbl_period.setText("Период не выбран")
        self.load_report_table()
    def load_all_tabs(self):
        self.main_window.db.update_overdue_fines()
        self.load_svodka()
        self.load_fines_table()
        self.load_report_table()
    def load_svodka(self):
        try:
            db = self.main_window.db
            db.connection.ping(reconnect=True)
            cur = db.connection.cursor()
            cur.execute("SELECT COUNT(*) AS c FROM traffic_fine")
            self.lbl_all.setText(f"Всего штрафов:\n{cur.fetchone()['c']}")
            cur.execute("SELECT COUNT(*) AS c FROM traffic_fine WHERE id_status_fine = 1")
            self.lbl_paid.setText(f"Оплаченных:\n{cur.fetchone()['c']}")
            cur.execute("SELECT COUNT(*) AS c FROM traffic_fine WHERE id_status_fine = 2")
            self.lbl_unpaid.setText(f"Не оплачено\n{cur.fetchone()['c']}")
            cur.execute("SELECT COUNT(*) AS c FROM traffic_fine WHERE id_status_fine = 3")
            self.lbl_debt.setText(f"Задолженность\n{cur.fetchone()['c']}")
            cur.execute(
                "SELECT COALESCE(SUM(fs.summary), 0) AS s FROM traffic_fine tf "
                "JOIN fine_sum fs ON tf.id_fine_sum = fs.id WHERE tf.id_status_fine IN (2, 3)"
            )
            self.lbl_sum.setText(f"Общий долг:\n{cur.fetchone()['s']} руб.")
            cur.close()
            cur = db.get_cursor()
            self.fill_table(cur, self.table_svodka, self.BASE_QUERY + " ORDER BY tf.id DESC LIMIT 10")
            cur.close()
        except Exception as e:
            print(f"Ошибка загрузки сводки: {e}")
    def load_fines_table(self):
        s = self.inp_search.text().strip()
        q = self.BASE_QUERY
        p = []
        if s:
            q += " WHERE tf.num_post LIKE %s"
            p.append(f"%{s}%")
        q += " ORDER BY tf.date_time DESC"
        cur = self.main_window.db.get_cursor()
        self.fill_table(cur, self.table_fines, q, p)
        cur.close()
    def load_report_table(self):
        s = self.inp_report_post.text().strip()
        q = self.BASE_QUERY
        p = []
        conditions = []
        if self.report_date_from and self.report_date_to:
            conditions.append("tf.date_time BETWEEN %s AND %s")
            p.append(self.report_date_from.toString("yyyy-MM-dd"))
            p.append(self.report_date_to.toString("yyyy-MM-dd") + " 23:59:59")
        if s:
            conditions.append("tf.num_post LIKE %s")
            p.append(f"%{s}%")
        if conditions:
            q += " WHERE " + " AND ".join(conditions)
        q += " ORDER BY tf.date_time DESC"
        cur = self.main_window.db.get_cursor()
        self.fill_table(cur, self.table_report, q, p)
        cur.close()
    def fill_table(self, cur, table_obj, query, params=None):
        cur.execute(query, params or [])
        rows = cur.fetchall()
        table_obj.setRowCount(len(rows))
        for i, r in enumerate(rows):
            actual_sp = r['actual_speed'] if r['actual_speed'] else "-"
            limit_sp = r['fine_speedcol'] if r['fine_speedcol'] else "-"
            speed_str = f"{actual_sp} / {limit_sp}"
            vals = [
                str(r['id']), str(r['date_time']), str(r['num_post']),
                str(r['fio']).strip(), str(r['car_num']), str(r['fine_place']),
                str(r.get('section_num') or '-'), speed_str,
                str(r['summary']), str(r['status']),
            ]
            for j, v in enumerate(vals):
                item = QTableWidgetItem(v)
                item.setTextAlignment(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
                table_obj.setItem(i, j, item)
        table_obj.resizeRowsToContents()
    def add_fine(self):
        if FineEditDialog(self.main_window).exec():
            self.load_all_tabs()
    def edit_fine(self):
        row = self.table_fines.currentRow()
        if row == -1:
            QMessageBox.warning(self, "Ошибка", "Выберите запись")
            return
        fid = self.table_fines.item(row, 0).text()
        if FineEditDialog(self.main_window, fine_id=fid).exec():
            self.load_all_tabs()
    def del_fine(self):
        row = self.table_fines.currentRow()
        if row == -1:
            return
        fid = self.table_fines.item(row, 0).text()
        if QMessageBox.question(
            self, "Подтверждение", f"Удалить постановление ID={fid}?"
        ) == QMessageBox.StandardButton.Yes:
            cur = self.main_window.db.get_cursor()
            cur.execute("DELETE FROM traffic_fine WHERE id = %s", (fid,))
            self.main_window.db.connection.commit()
            cur.close()
            self.load_all_tabs()
    def export_excel(self):
        if not self.report_date_from or not self.report_date_to:
            QMessageBox.warning(self, "Внимание", "Сначала выберите период")
            return
        date_from = self.report_date_from.toString("yyyy-MM-dd")
        date_to = self.report_date_to.toString("yyyy-MM-dd") + " 23:59:59"
        try:
            cur = self.main_window.db.get_cursor()
            cur.execute(
                self.BASE_QUERY
                + " WHERE tf.date_time BETWEEN %s AND %s"
                + " AND tf.id_status_fine IN (2, 3)"
                + " ORDER BY tf.date_time DESC",
                (date_from, date_to)
            )
            rows = cur.fetchall()
            cur.close()
        except Exception as e:
            QMessageBox.critical(self, "Ошибка БД", str(e))
            return
        if not rows:
            QMessageBox.information(self, "Нет данных", "За выбранный период неоплаченных постановлений не найдено.")
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Постановления"
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        col_headers = ["ID", "Дата", "Постановление", "ФИО", "Госномер",
                       "Место", "Статья", "Факт / Лимит ск.", "Сумма", "Статус"]
        col_widths = [6, 18, 20, 24, 12, 30, 16, 16, 10, 14]
        for col_idx, (header, width) in enumerate(zip(col_headers, col_widths), start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(name="Arial", bold=True, size=11)
            cell.alignment = center
            cell.border = border
            ws.column_dimensions[cell.column_letter].width = width
        ws.row_dimensions[1].height = 18
        for row_idx, r in enumerate(rows, start=2):
            actual_sp = r['actual_speed'] if r['actual_speed'] else "-"
            limit_sp = r['fine_speedcol'] if r['fine_speedcol'] else "-"
            speed_str = f"{actual_sp} / {limit_sp}"
            values = [
                r['id'], str(r['date_time']), r['num_post'],
                str(r['fio']).strip(), r['car_num'], r['fine_place'],
                str(r.get('section_num') or '-'), speed_str,
                r['summary'], r['status'],
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = Font(name="Arial", size=10)
                cell.border = border
                cell.alignment = left if col_idx in (4, 6) else center
        safe_from = self.report_date_from.toString("yyyyMMdd")
        safe_to = self.report_date_to.toString("yyyyMMdd")
        filename = f"report_{safe_from}_{safe_to}.xlsx"
        filepath = os.path.join(os.path.dirname(os.path.abspath(__file__)), filename)
        try:
            wb.save(filepath)
            QMessageBox.information(self, "Готово", f"Отчёт сохранён:\n{filepath}")
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить файл:\n{e}")
    def export_word(self):
        row = self.table_report.currentRow()
        if row == -1:
            QMessageBox.warning(self, "Ошибка", "Выберите строку в таблице отчёта")
            return
        fine_id = self.table_report.item(row, 0).text()
        num_post = self.table_report.item(row, 2).text()
        try:
            cur = self.main_window.db.get_cursor()
            cur.execute("""
                SELECT d.surname AS drv_surname, d.name AS drv_name,
                       d.middle_name AS drv_middle, d.place_of_birth AS drv_birth,
                       d.date_of_birth AS drv_dob, d.id_gender,
                       c.car_model, c.car_num,
                       tf.date_time, tf.num_post, tf.fine_place, fs.summary,
                       tf.actual_speed, fsp.fine_speedcol, fsec.section_num,
                       TRIM(CONCAT(e.surname,' ',e.name,' ',COALESCE(e.middle_name,''))) AS emp_fio
                FROM traffic_fine tf
                JOIN car c ON tf.id_car = c.id
                JOIN driver d ON c.id_driver = d.id
                JOIN fine_sum fs ON tf.id_fine_sum = fs.id
                JOIN emp_gibdd e ON tf.id_emp = e.id
                LEFT JOIN fine_speed fsp ON tf.id_speed_fine = fsp.id
                LEFT JOIN fine_section fsec ON tf.id_fine_section = fsec.id
                WHERE tf.id = %s
            """, (fine_id,))
            rec = cur.fetchone()
            cur.close()
        except Exception as e:
            QMessageBox.critical(self, "Ошибка БД", str(e))
            return
        if not rec:
            return
        dt_obj = rec['date_time']
        if isinstance(dt_obj, datetime.datetime):
            date_str = dt_obj.strftime("%d.%m.%Y")
            time_str = dt_obj.strftime("%H:%M")
        else:
            parts = str(dt_obj).split()
            date_str = parts[0] if parts else str(dt_obj)
            time_str = parts[1] if len(parts) > 1 else ""
        drv_fio = f"{rec['drv_surname']} {rec['drv_name']} {rec.get('drv_middle') or ''}".strip()
        fine_sum = rec['summary']
        half_sum = round(float(fine_sum) / 2, 2) if fine_sum else "—"
        actual_sp = rec.get('actual_speed')
        limit_sp = rec.get('fine_speedcol')
        actual_speed_str = str(actual_sp) if actual_sp else ""
        limit_speed_str = str(limit_sp) if limit_sp else ""
        up_speed_str = (str(actual_sp - limit_sp) if (actual_sp and limit_sp and actual_sp > limit_sp) else "0")
        section_num = str(rec.get('section_num') or "")
        replacements = {
            "{{ id }}": num_post, "{{ date }}": date_str, "{{ time }}": time_str,
            "{{ address }}": str(rec['fine_place']), "{{ car_mark }}": str(rec['car_model']),
            "{{ car_num }}": str(rec['car_num']),
            "{{ actual_speed_fine }}": actual_speed_str,
            "{{ speed_fine }}": limit_speed_str, "{{ up_speed_car }}": up_speed_str,
            "{{ car_id }}": str(rec['car_num']), "{{ FIO_driver }}": drv_fio,
            "{{ date_of_birth }}": str(rec.get('drv_dob') or ""),
            "{{ place_of_birth }}": str(rec.get('drv_birth') or ""),
            "{{ section_num }}": section_num, "{{ fine_sum }}": str(fine_sum),
            "{{ half_of_fine_sum }}": str(half_sum),
            "{{ employee }}": str((rec['emp_fio'] or "").strip()),
        }
        template_file = "shablon1.docx" if rec.get('id_gender') == 1 else "shablon2.docx"
        def replace_in_paragraph(paragraph, rd):
            full_text = "".join(r.text for r in paragraph.runs)
            changed = False
            for key, value in rd.items():
                if key in full_text:
                    full_text = full_text.replace(key, value)
                    changed = True
            if changed and paragraph.runs:
                paragraph.runs[0].text = full_text
                for run in paragraph.runs[1:]:
                    run.text = ""
        try:
            template_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), template_file)
            if not os.path.exists(template_path):
                QMessageBox.critical(self, "Ошибка", f"Файл шаблона не найден:\n{template_path}")
                return
            doc = docx.Document(template_path)
            for paragraph in doc.paragraphs:
                replace_in_paragraph(paragraph, replacements)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            replace_in_paragraph(paragraph, replacements)
            safe_num = num_post.replace("/", "-").replace("\\", "-")
            filename = f"postanovlenie_{safe_num}.docx"
            doc.save(filename)
            QMessageBox.information(self, "Готово", f"Документ «{filename}» успешно сформирован")
        except Exception as e:
            import traceback
            traceback.print_exc()
            QMessageBox.critical(self, "Ошибка", f"Не удалось создать Word-файл:\n{e}")
    def logout(self):
        self.main_window.emp_data = None
        self.main_window.stacked.setCurrentIndex(0)
class FineEditDialog(QDialog):
    def __init__(self, main_window, fine_id=None):
        super().__init__()
        self.main_window = main_window
        self.fine_id = fine_id
        self.build_ui()
        self.load_dictionaries()
        if self.fine_id:
            self.load_fine()
        else:
            self.inp_date.setText(datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    def build_ui(self):
        self.setWindowTitle("Редактирование штрафа" if self.fine_id else "Добавление штрафа")
        self.resize(480, 580)
        layout = QVBoxLayout(self)
        layout.setSpacing(8)
        self.inp_date = QLineEdit()
        self.inp_date.setPlaceholderText("Дата (ГГГГ-ММ-ДД ЧЧ:ММ:СС)")
        self.inp_plate = QLineEdit()
        self.inp_plate.setPlaceholderText("Госномер автомобиля")
        self.inp_place = QLineEdit()
        self.inp_place.setPlaceholderText("Место нарушения")
        self.combo_section = QComboBox()
        self.combo_speed_limit = QComboBox()
        self.combo_speed_limit.addItem("Без нарушения скорости", None)
        self.inp_actual_speed = QLineEdit()
        self.inp_actual_speed.setPlaceholderText("Фактическая скорость (только цифры)")
        self.combo_sum = QComboBox()
        self.combo_stat = QComboBox()
        def add_labeled(text, widget):
            lbl = QLabel(text)
            lbl.setFont(QFont("Arial", 11))
            widget.setFixedSize(450, 40)
            widget.setFont(QFont("Arial", 12))
            layout.addWidget(lbl)
            layout.addWidget(widget, alignment=Qt.AlignmentFlag.AlignHCenter)
        for w in [self.inp_date, self.inp_plate, self.inp_place]:
            w.setFixedSize(450, 40)
            w.setFont(QFont("Arial", 12))
            layout.addWidget(w, alignment=Qt.AlignmentFlag.AlignHCenter)
        add_labeled("Статья КоАП:", self.combo_section)
        add_labeled("Ограничение скорости на участке:", self.combo_speed_limit)
        self.inp_actual_speed.setFixedSize(450, 40)
        self.inp_actual_speed.setFont(QFont("Arial", 12))
        layout.addWidget(self.inp_actual_speed, alignment=Qt.AlignmentFlag.AlignHCenter)
        add_labeled("Сумма штрафа:", self.combo_sum)
        self.lbl_stat = QLabel("Статус:")
        self.lbl_stat.setFont(QFont("Arial", 11))
        self.combo_stat.setFixedSize(450, 40)
        self.combo_stat.setFont(QFont("Arial", 12))
        layout.addWidget(self.lbl_stat)
        layout.addWidget(self.combo_stat, alignment=Qt.AlignmentFlag.AlignHCenter)
        self.lbl_stat.setVisible(bool(self.fine_id))
        self.combo_stat.setVisible(bool(self.fine_id))
        btn_save = QPushButton("Сохранить")
        btn_save.setFixedSize(300, 45)
        btn_save.setFont(QFont("Arial", 14))
        btn_save.clicked.connect(self.save_data)
        layout.addWidget(btn_save, alignment=Qt.AlignmentFlag.AlignHCenter)
    def load_dictionaries(self):
        try:
            cur = self.main_window.db.get_cursor()
            cur.execute("SELECT id, summary FROM fine_sum ORDER BY summary")
            for r in cur.fetchall():
                self.combo_sum.addItem(f"{r['summary']} руб.", r['id'])
            cur.execute("SELECT id, name FROM status_fine ORDER BY id")
            for r in cur.fetchall():
                self.combo_stat.addItem(str(r['name']), r['id'])
            cur.execute("SELECT id, section_num FROM fine_section ORDER BY id")
            for r in cur.fetchall():
                self.combo_section.addItem(str(r['section_num']), r['id'])
            cur.execute("SELECT id, fine_speedcol FROM fine_speed ORDER BY fine_speedcol")
            for r in cur.fetchall():
                self.combo_speed_limit.addItem(f"{r['fine_speedcol']} км/ч", r['id'])
            cur.close()
        except Exception as e:
            print(f"Ошибка комбобоксов: {e}")
    def load_fine(self):
        try:
            cur = self.main_window.db.get_cursor()
            cur.execute(
                "SELECT tf.date_time, c.car_num, tf.fine_place, "
                "tf.id_fine_sum, tf.id_status_fine, "
                "tf.id_fine_section, tf.id_speed_fine, tf.actual_speed "
                "FROM traffic_fine tf JOIN car c ON tf.id_car = c.id "
                "WHERE tf.id = %s", (self.fine_id,)
            )
            f = cur.fetchone()
            cur.close()
            if f:
                self.inp_date.setText(str(f['date_time']))
                self.inp_plate.setText(str(f['car_num']))
                self.inp_place.setText(str(f['fine_place']))
                self.combo_sum.setCurrentIndex(self.combo_sum.findData(f['id_fine_sum']))
                self.combo_stat.setCurrentIndex(self.combo_stat.findData(f['id_status_fine']))
                if f['id_fine_section']:
                    self.combo_section.setCurrentIndex(self.combo_section.findData(f['id_fine_section']))
                if f['id_speed_fine']:
                    self.combo_speed_limit.setCurrentIndex(self.combo_speed_limit.findData(f['id_speed_fine']))
                if f['actual_speed']:
                    self.inp_actual_speed.setText(str(f['actual_speed']))
        except Exception as e:
            print(f"Ошибка данных штрафа: {e}")
    def save_data(self):
        val_date = self.inp_date.text().strip()
        val_plate = self.inp_plate.text().strip().upper()
        val_place = self.inp_place.text().strip()
        val_sum = self.combo_sum.currentData()
        val_section = self.combo_section.currentData()
        val_speed_limit = self.combo_speed_limit.currentData()
        val_actual_speed = self.inp_actual_speed.text().strip()
        if not all([val_date, val_plate, val_place]):
            QMessageBox.warning(self, "Ошибка", "Заполните дату, госномер и место")
            return
        if val_actual_speed and not val_actual_speed.isdigit():
            QMessageBox.warning(self, "Ошибка", "Фактическая скорость должна быть числом")
            return
        val_actual_speed = int(val_actual_speed) if val_actual_speed else None
        try:
            cur = self.main_window.db.get_cursor()
            cur.execute("SELECT id FROM car WHERE car_num = %s", (val_plate,))
            car_row = cur.fetchone()
            if not car_row:
                QMessageBox.warning(self, "Ошибка", "Госномер не зарегистрирован в системе")
                cur.close()
                return
            car_id = car_row['id']
            emp_id = self.main_window.emp_data['id']
            if self.fine_id:
                val_stat = self.combo_stat.currentData()
                cur.execute(
                    "UPDATE traffic_fine "
                    "SET date_time=%s, id_car=%s, fine_place=%s, "
                    "id_fine_sum=%s, id_status_fine=%s, "
                    "id_fine_section=%s, id_speed_fine=%s, actual_speed=%s "
                    "WHERE id=%s",
                    (val_date, car_id, val_place, val_sum, val_stat,
                     val_section, val_speed_limit, val_actual_speed, self.fine_id)
                )
            else:
                val_post = "П-" + datetime.datetime.now().strftime("%Y%m%d%H%M%S")
                cur.execute(
                    "INSERT INTO traffic_fine "
                    "(id_emp, id_car, date_time, num_post, fine_place, "
                    "id_fine_sum, id_status_fine, id_fine_section, id_speed_fine, actual_speed) "
                    "VALUES (%s, %s, %s, %s, %s, %s, 2, %s, %s, %s)",
                    (emp_id, car_id, val_date, val_post, val_place,
                     val_sum, val_section, val_speed_limit, val_actual_speed)
                )
            self.main_window.db.connection.commit()
            cur.close()
            self.accept()
        except Exception as e:
            self.main_window.db.connection.rollback()
            QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить:\n{e}")
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("ИС мониторинга штрафов ГИБДД")
        self.resize(1280, 800)
        self.db = DatabaseManager()
        self.driver_data = None
        self.emp_data = None
        self.stacked = QStackedWidget()
        self.setCentralWidget(self.stacked)
        self.window_role = RoleWindow(self)
        self.window_driver_login = DriverLoginWindow(self)
        self.window_driver_reg = DriverRegWindow(self)
        self.window_driver_main = DriverMainWindow(self)
        self.window_driver_pay = DriverPayWindow(self)
        self.window_emp_login = EmployeeLoginWindow(self)
        self.window_emp_main = EmployeeMainWindow(self)
        for w in [self.window_role, self.window_driver_login, self.window_driver_reg,
                  self.window_driver_main, self.window_driver_pay,
                  self.window_emp_login, self.window_emp_main]:
            self.stacked.addWidget(w)
        self.stacked.setCurrentIndex(0)
    def close_event(self, event):
        self.db.close_db()
        event.accept()
if __name__ == '__main__':
    try:
        app = QApplication(sys.argv)
        app.setStyle("Fusion")
        window = MainWindow()
        window.show()
        sys.exit(app.exec())
    except Exception as e:
        input(f"Критическая ошибка: {e}\nНажмите Enter...")
